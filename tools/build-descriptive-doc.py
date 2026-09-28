"""Builds design/Trust-No-Wall-Descriptive-Document.docx, the course's descriptive document.

Usage: python tools/build-descriptive-doc.py [VIDEO_URL]
Requires python-docx. Content mirrors design/DesignDocument.md and planning/spec.md.
"""
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "design" / "Trust-No-Wall-Descriptive-Document.docx"
REPO = "https://github.com/CSCI-526/Team7-trust-no-walls"
PLAY = "https://csci-526.github.io/Team7-trust-no-walls/"
VIDEO = sys.argv[1] if len(sys.argv) > 1 else "https://drive.google.com/file/d/108_E55riLKDuFTiWyroH48iHKu0jDZD_/view?usp=sharing"

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
for section in doc.sections:
    section.left_margin = section.right_margin = Inches(0.8)
    section.top_margin = section.bottom_margin = Inches(0.8)


def heading(text, level=1):
    doc.add_heading(text, level=level)


def para(text="", bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
    p.add_run(text)
    return p


def bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
    p.add_run(text)
    return p


def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def table(header, rows, widths=None, font_size=9):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(font_size)
        shade(c, "D9E2F3")
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(val)
            r.font.size = Pt(font_size)
            if i == 0:
                r.bold = True
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return t


# Title
title = doc.add_heading("Trust No Wall", level=0)
p = doc.add_paragraph()
p.add_run("Descriptive Document | CSCI 526 Paired Prototype | Team 7").italic = True
para("Qiming Xiao, Manas Vardhan", bold_prefix="Team: ")
para(PLAY, bold_prefix="Play online: ")
para(VIDEO, bold_prefix="Gameplay video: ")
para(REPO, bold_prefix="Repository: ")

# Logline
heading("Logline (Genre + Twist)")
para("A 2D top-down maze game where the maze itself lies to you: walls hide, move, vanish and rewire, "
     "but every lie follows a learnable rule (Maze + Deceptive, Changing Maze).")

# Research
heading("Genre Tropes Research and Twist")
heading("Research: three games from the maze genre", level=2)
para("We looked at three maze games spanning the genre from the arcade era to modern browser games: "
     "Pac-Man, Maze Craze: A Game of Cops 'n Robbers, and Level Devil.")
bullet(" (Namco, arcade, 1980). The player navigates a fixed, fully visible maze, collecting pellets while avoiding "
       "four ghosts, each with its own movement pattern. The maze never changes: what you see is exactly what is there, and "
       "mastery comes from learning the ghosts' movement, not the maze.", bold_prefix="Pac-Man")
bullet(" (Atari, Atari 2600). Two players race through a randomly generated top-down maze. Some optional game "
       "variants add novelty rules, such as a maze that is invisible and only briefly revealed, or a false wall a "
       "player can drop to block an opponent. These are side modes layered on an otherwise honest maze.",
       bold_prefix="Maze Craze: A Game of Cops 'n Robbers")
bullet(" (Unept (Adam Corey), browser 2023). A minimalist trial-and-error game in which floors turn into spikes, "
       "ceilings fall and platforms vanish the moment you commit. The level actively deceives the player, who "
       "progresses only by dying, memorizing and retrying the identical level, with no telegraphing on a first "
       "attempt.", bold_prefix="Level Devil")

heading("Shared genre tropes", level=2)
bullet("The maze is a stable, observable puzzle: what you can see is the truth of the space, and skill means reading "
       "the visible layout correctly.")
bullet("Hidden or false information, when it exists (Maze Craze's invisible-maze variant), is an optional novelty "
       "mode, not the core condition of play.")
bullet("When deception is the whole point (Level Devil), it has no fairness guarantee: nothing checks that a level "
       "stays completable, and a trap can be an unavoidable first-try loss with no rule to learn from.")
bullet("Failing sends the player back to the same, unchanged challenge: only the player's knowledge improves.")
bullet("Simple directional movement, a clear start and goal, and difficulty that rises through larger or more "
       "complex layouts and more threats.")

heading("Twist and why it is innovative", level=2)
para("Every researched game assumes a static, trustworthy maze solved by observation. Trust No Wall breaks that "
     "assumption: the maze is deceptive (invisible walls, one-way corridors that look open both ways, a decoy goal) "
     "and reactive (walls that slide, vanish and rewire on triggers, floors that collapse, and a shadow that follows "
     "your own trail). Two guarantees keep it a fair maze game rather than a troll game: every generated level is "
     "checked by a solvability validator before it is shown, so a completable route always exists and no reachable "
     "cell is a trap; and every lie follows a fixed, telegraphed rule (a wall flashes and shows an arrow before it "
     "slides, plates share a color with the walls they control, memory tiles reveal nearby hidden walls, the real "
     "goal's star spins while a decoy's does not).", bold_prefix="Twist: ")
para("", bold_prefix="Why it is innovative: ")
bullet("It subverts the honest, fully visible maze (Pac-Man) and turns hidden information from an occasional side "
       "mode (Maze Craze) into the constant condition of the whole game.", bold_prefix="Subversion. ")
bullet("It combines Maze Craze's hidden-maze idea with Level Devil's committed deception, but applies it to spatial "
       "maze navigation instead of platforming reflexes.", bold_prefix="Combination. ")
bullet("It fixes Level Devil's weak point with a solvability guarantee and learnable, telegraphed rules for every "
       "hazard, so deception is always fair.", bold_prefix="Extension. ")
bullet("It keeps the restart-to-the-same-challenge trope on purpose: dying restarts the same layout, so a lie you "
       "die to once becomes a rule you can plan around, and remembering the maze's specific lies becomes the skill.",
       bold_prefix="Emphasis. ")

# Prototype description
heading("Short Prototype Description")
para("The player steps through a top-down grid maze from a Start pad to a Destination star, and the maze grows by one "
     "cell in each dimension every level (4 x 4 on level 1). As levels progress, new obstacles that deceive or change the "
     "maze unlock: invisible walls, sliding and disappearing walls, collapsing floors, trigger plates that rewire walls, "
     "teleporters, rotating gates, patrols, one-way paths, a decoy goal and a shadow that follows the player's trail. "
     "Because each obstacle follows a fixed rule and dying restarts the same layout, the player must observe, time and "
     "remember the maze's tricks instead of simply reading its map.")

heading("Controls", level=2)
table(["Input", "Action"], [
    ["WASD / Arrow keys", "Move one cell per step; hold to keep moving"],
    ["Space / Enter", "Start the game; skip an intro card"],
    ["R", "Restart the current attempt (same layout)"],
    ["P / Esc", "Pause"],
    ["M", "Mute"],
    ["F2", "Toggle the demo autopilot"],
], widths=[1.8, 4.9], font_size=10)

# Mechanics matrix
heading("Mechanic Matrices: Twist and Mechanics Matrix (core mechanic)")
matrix = [
    ["Deceptive, Changing Maze (Core Mechanic)",
     "The maze's walls and floor hide, move, vanish and rewire according to fixed rules while the player navigates it.",
     "This is the twist: the player cannot trust what they see and must learn each rule instead of just reading the map.",
     "Maze layout, walls, navigation, goal", "Subversion, Combination", "Deception; learnable rules"],
]
table(["Mechanics", "Description", "Interaction with Twist", "Affected Genre Elements", "Type of Genre Innovation", "Supports"],
      matrix, widths=[1.1, 1.6, 1.6, 0.9, 0.8, 0.8], font_size=9)

heading("Level progression", level=2)
table(["Level", "Maze", "New mechanics"], [
    ["1", "4 x 4", "None: basic maze with Start and Destination"],
    ["2", "5 x 5", "Invisible Walls, Memory Tiles"],
    ["3", "6 x 6", "Moving Walls, Disappearing Walls"],
    ["4", "7 x 7", "Collapsing Tiles, Trigger Walls (+ one earlier mechanic)"],
    ["5", "8 x 8", "Teleport Tiles, Rotating Barriers (+ one earlier mechanic)"],
    ["6", "9 x 9", "Patrolling Obstacles, One-Way Paths (+ one earlier mechanic)"],
    ["7", "10 x 10", "Decoy Destination (+ one earlier mechanic)"],
    ["8", "11 x 11", "Chasing Hazard (+ one earlier mechanic)"],
    ["9+", "12 x 12 and up", "4 random mechanics per level (5 from level 12)"],
], widths=[0.8, 1.3, 4.6], font_size=10)

# Repo
heading("GitHub Repository")
para(REPO, bold_prefix="Repository: ")
para(PLAY, bold_prefix="Playable build (GitHub Pages): ")

# Contributions
heading("Individual Contributions")
bullet("Game concept and design direction (maze + deceptive-maze twist), genre research on Pac-Man, Maze Craze and Level Devil, design of the obstacle set and level progression (which mechanic unlocks when), playtesting and difficulty feedback, the descriptive document, and the gameplay video.", bold_prefix="Qiming Xiao: ")
bullet("Unity implementation: procedural maze generation and the solvability validator, the level simulation for all obstacles, rendering, animations, HUD, audio and game flow, the demo autopilot and automated tests, and the WebGL build hosted on GitHub Pages.", bold_prefix="Manas Vardhan: ")

# Diagrams
heading("Sketches and Diagrams")
para("Game loop: from the title screen through intro cards, play, death (restart the same layout) and level "
     "completion (next, larger maze).", bold_prefix="Figure 1. ")
doc.add_picture(str(ROOT / "design" / "diagrams" / "game-loop.png"), width=Inches(6.6))
para("Level generation: a seeded maze is generated, loops are added, mechanics are chosen by level, and every "
     "obstacle is placed only if the solvability validator confirms the level stays completable with no traps.",
     bold_prefix="Figure 2. ")
doc.add_picture(str(ROOT / "design" / "diagrams" / "level-generation.png"), width=Inches(6.6))

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(f"wrote {OUT}")
